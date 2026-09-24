"""Gera artigo/rascunho-artigo.docx a partir de artigo/rascunho-artigo.md.

Usa o pandoc (pacote pypandoc_binary) com o modelo artigo/modelo/referencia-abnt.docx
(Times New Roman 12, espaçamento simples, margens 3/3/2/2 cm) e depois ajusta as tabelas
com python-docx: fonte 10, tabelas abertas no padrão AER/Springer (filete em cima, sob o cabeçalho e embaixo, sem
grade), rótulo do título em negrito e "Fonte" em itálico. Caminhos de dados viram links para o GitHub e DOIs para
doi.org (artigo/links.py). Com --pdf, converte
também para PDF pelo LibreOffice, só para conferir o número de páginas (10 a 15).

Uso: python artigo/gerar_docx.py [--pdf]
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from links import codigo_em_links_md, dois_em_links

ART = Path(__file__).resolve().parent
MD = ART / "rascunho-artigo.md"
DOCX = ART / "rascunho-artigo.docx"
MODELO = ART / "modelo" / "referencia-abnt.docx"


def _borda(pai, lado: str, sz: str) -> None:
    e = OxmlElement(f"w:{lado}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), sz)
    e.set(qn("w:space"), "0")
    e.set(qn("w:color"), "000000")
    pai.append(e)


def bordas(tabela) -> None:
    """Tabela aberta (AER/Springer): filete de 1 pt em cima e embaixo, 0,5 pt sob o cabeçalho, sem grade."""
    tblPr = tabela._tbl.tblPr
    for velho in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(velho)
    b = OxmlElement("w:tblBorders")
    _borda(b, "top", "8")
    _borda(b, "bottom", "8")
    for lado in ("left", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)
    for cel in tabela.rows[0].cells:
        tcPr = cel._tc.get_or_add_tcPr()
        tcb = OxmlElement("w:tcBorders")
        _borda(tcb, "bottom", "4")
        tcPr.append(tcb)


def ajustar_tabelas(caminho: Path) -> None:
    doc = Document(caminho)
    # 4 pt entre parágrafos (o mesmo da versão LaTeX); espaçamento simples mantido
    for nome in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if nome in [s.name for s in doc.styles]:
            doc.styles[nome].paragraph_format.space_after = Pt(4)
    # título de cada tabela = parágrafo imediatamente anterior no corpo do documento
    titulos = []  # na mesma ordem de doc.tables (tabelas de nível superior, em ordem do documento)
    anterior = ""
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            anterior = "".join(x.text or "" for x in el.iter(qn("w:t")))
        elif el.tag == qn("w:tbl"):
            titulos.append(anterior)
    for t, titulo in zip(doc.tables, titulos):
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        bordas(t)
        for i, linha in enumerate(t.rows):
            for cel in linha.cells:
                for par in cel.paragraphs:
                    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    par.paragraph_format.space_after = Pt(0)
                    par.paragraph_format.space_before = Pt(0)
                    for run in par.runs:
                        run.font.size = Pt(10)
        # cabeçalho repetido em quebra de página
        cab = t.rows[0]._tr.get_or_add_trPr()
        rep = OxmlElement("w:tblHeader")
        rep.set(qn("w:val"), "true")
        cab.append(rep)
        # só tabelas numéricas curtas ficam inteiras; quadros (texto) podem se dividir entre páginas
        if titulo.startswith("Tabela") and len(t.rows) <= 8:
            for linha in t.rows[:-1]:
                for cel in linha.cells:
                    for par in cel.paragraphs:
                        par.paragraph_format.keep_with_next = True
    # título de tabela, quadro e figura: mesma página do objeto; só o rótulo em negrito
    for par in doc.paragraphs:
        m = re.match(r"((?:Tabela|Quadro|Figura) \d+)(.*)", par.text)
        if m:
            par.paragraph_format.keep_with_next = True
            for run in list(par.runs):
                run._r.getparent().remove(run._r)
            par.add_run(m.group(1)).font.bold = True
            par.add_run(m.group(2)).font.bold = False
    # "Fonte:" logo abaixo de tabelas e figuras, em corpo 10, com o rótulo em itálico
    for par in doc.paragraphs:
        if par.text.startswith("Fonte:") and par.runs:
            primeiro = par.runs[0]
            if primeiro.text.startswith("Fonte"):
                primeiro.text = primeiro.text[len("Fonte"):]
                rotulo = par.add_run("Fonte")
                rotulo.italic = True
                primeiro._r.addprevious(rotulo._r)
            for run in par.runs:
                run.font.size = Pt(10)
    doc.save(caminho)


def main() -> None:
    texto = codigo_em_links_md(dois_em_links(MD.read_text(encoding="utf-8")))
    pypandoc.convert_text(
        texto, "docx", format="markdown", outputfile=str(DOCX),
        extra_args=[f"--reference-doc={MODELO}", f"--resource-path={ART}", "--wrap=none"],
    )
    ajustar_tabelas(DOCX)
    print(DOCX)
    if "--pdf" in sys.argv:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(ART), str(DOCX)],
                       check=True, capture_output=True)
        print(ART / "rascunho-artigo.pdf")


if __name__ == "__main__":
    main()
